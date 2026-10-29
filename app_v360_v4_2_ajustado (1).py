from flask import (
    Flask,
    request,
    redirect,
    flash,
    render_template_string
)

from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime
import sqlite3
import re


# ============================================================
# CONFIGURAÇÃO
# ============================================================

app = Flask(__name__)
app.secret_key = "v360-v42"

ROOT = Path(__file__).parent
DB = ROOT / "v360.db"

UFS = (
    "CE MA PI TO PB RN SE AL BA PE MT AM PA RR RO AC AP"
).split()

SECS = {
    "SSD",
    "STILL",
    "CERVEJA",
    "NOVOS NEGÓCIOS"
}


# ============================================================
# BANCO DE DADOS
# ============================================================

def cn():
    conexao = sqlite3.connect(DB, timeout=20)
    conexao.row_factory = sqlite3.Row
    return conexao


def init():

    with cn() as c:

        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS products(
                id INTEGER PRIMARY KEY,
                state TEXT,
                category TEXT,
                product TEXT,
                base_factor INTEGER,
                matrix TEXT
            );

            CREATE TABLE IF NOT EXISTS requests(
                id INTEGER PRIMARY KEY,
                created_at TEXT,
                requester TEXT,
                area TEXT,
                product_id INTEGER,
                state TEXT,
                category TEXT,
                product TEXT,
                old_factor INTEGER,
                new_factor INTEGER,
                competence TEXT,
                justification TEXT,
                status TEXT DEFAULT 'Pendente',
                decided_at TEXT,
                decided_by TEXT,
                decision_reason TEXT
            );

            CREATE TABLE IF NOT EXISTS periods(
                id INTEGER PRIMARY KEY,
                request_id INTEGER,
                product_id INTEGER,
                competence TEXT,
                new_factor INTEGER,
                approved_at TEXT,
                approved_by TEXT,
                UNIQUE(product_id, competence)
            );

            CREATE TABLE IF NOT EXISTS audit(
                id INTEGER PRIMARY KEY,
                created_at TEXT,
                action TEXT,
                details TEXT
            );
            """
        )


def log(action, details):

    with cn() as c:

        c.execute(
            """
            INSERT INTO audit(
                created_at,
                action,
                details
            )
            VALUES(?,?,?)
            """,
            (
                datetime.now().isoformat(timespec="seconds"),
                action,
                details
            )
        )


# ============================================================
# LEITURA DA MATRIZ
# ============================================================

def extract(path):

    workbook = load_workbook(
        path,
        data_only=True
    )

    output = []

    for worksheet in workbook.worksheets:

        label = (
            str(worksheet.cell(2, 2).value or "")
            + " "
            + worksheet.title
        ).upper()

        states = list(
            dict.fromkeys(
                re.findall(
                    r"\b(?:" + "|".join(UFS) + r")\b",
                    label
                )
            )
        )

        section = None
        factor_columns = {}

        for row_number in range(
            1,
            worksheet.max_row + 1
        ):

            values = [
                worksheet.cell(
                    row_number,
                    column_number
                ).value
                for column_number in range(
                    1,
                    worksheet.max_column + 1
                )
            ]

            texts = [
                str(value).strip()
                for value in values
                if value is not None
            ]

            section_found = next(
                (
                    text.upper()
                    for text in texts
                    if text.upper() in SECS
                ),
                None
            )

            if section_found:

                section = (
                    section_found.title()
                    if section_found != "SSD"
                    else "SSD"
                )

                factor_columns = {}
                continue

            if not section:
                continue

            for column_number, value in enumerate(
                values,
                1
            ):

                match = re.search(
                    r"Fator\s*(\d+)",
                    str(value or ""),
                    re.I
                )

                if match:

                    factor_columns[column_number] = int(
                        match.group(1)
                    )

            if any(
                str(value or "").strip().lower()
                == "produto"
                for value in values
            ):
                continue

            for column_number, factor_value in (
                factor_columns.items()
            ):

                product_name = str(
                    worksheet.cell(
                        row_number,
                        column_number
                    ).value
                    or ""
                ).strip()

                if (
                    product_name
                    and product_name.upper() not in SECS
                    and "GRUPO"
                    not in product_name.upper()
                    and not product_name.startswith("*")
                ):

                    for state in states:

                        output.append(
                            (
                                state,
                                section,
                                product_name,
                                factor_value,
                                worksheet.title
                            )
                        )

    return output


# ============================================================
# FATOR POR COMPETÊNCIA
# ============================================================

def factor(connection, product_id, competence):

    result = connection.execute(
        """
        SELECT new_factor
        FROM periods
        WHERE product_id = ?
          AND competence <= ?
        ORDER BY competence DESC, id DESC
        LIMIT 1
        """,
        (
            product_id,
            competence
        )
    ).fetchone()

    if result:
        return result["new_factor"]

    product = connection.execute(
        """
        SELECT base_factor
        FROM products
        WHERE id = ?
        """,
        (product_id,)
    ).fetchone()

    return product["base_factor"]


def nowcomp():
    return datetime.now().strftime("%Y-%m")


# ============================================================
# FORMATAÇÃO DA COMPETÊNCIA
# ============================================================

def lab(value):

    if value is None:
        return "Competência não informada"

    value = str(value).strip()

    if not value:
        return "Competência não informada"

    try:

        year, month = value.split("-")

        months = [
            "",
            "Jan",
            "Fev",
            "Mar",
            "Abr",
            "Mai",
            "Jun",
            "Jul",
            "Ago",
            "Set",
            "Out",
            "Nov",
            "Dez"
        ]

        month_number = int(month)

        if month_number < 1 or month_number > 12:
            return value

        return f"{months[month_number]}/{year}"

    except (ValueError, TypeError, IndexError):
        return value


app.jinja_env.filters["comp"] = lab


# ============================================================
# CSS
# ============================================================

CSS = """
<style>

:root {
    --wine: #7b1233;
    --background: #f3f5f7;
    --line: #e1e5e9;
    --muted: #687386;
}

* {
    box-sizing: border-box;
}

body {
    margin: 0;
    font-family: "Segoe UI", Arial, sans-serif;
    background: var(--background);
    color: #111827;
}

aside {
    position: fixed;
    inset: 0 auto 0 0;
    width: 258px;
    background: linear-gradient(
        165deg,
        #171318,
        #68112e
    );
    color: #ffffff;
    padding: 24px 18px;
}

.logo {
    font-size: 29px;
    font-weight: 900;
}

.logo b {
    background: #ffffff;
    color: var(--wine);
    padding: 4px 10px;
    border-radius: 10px;
}

.logo small {
    display: block;
    font-size: 11px;
    letter-spacing: 2px;
    opacity: 0.7;
    margin-top: 9px;
}

nav {
    margin-top: 38px;
}

nav a {
    display: block;
    color: #ffffff;
    text-decoration: none;
    padding: 12px;
    border-radius: 9px;
}

nav a:hover {
    background: #ffffff18;
}

.main {
    margin-left: 258px;
}

.top {
    background: #ffffff;
    height: 74px;
    padding: 17px 32px;
    border-bottom: 1px solid var(--line);
}

.top span {
    display: block;
    color: var(--muted);
    font-size: 12px;
}

.content {
    padding: 28px 35px;
}

.panel,
.card {
    background: #ffffff;
    border: 1px solid var(--line);
    border-radius: 16px;
    box-shadow: 0 8px 24px #00000008;
}

.panel {
    padding: 21px;
    margin-top: 18px;
}

.grid {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 14px;
}

.card {
    padding: 19px;
}

.card strong {
    display: block;
    font-size: 29px;
    margin-top: 7px;
}

.title h1 {
    margin: 0;
}

.title p {
    color: var(--muted);
}

.toolbar {
    display: flex;
    gap: 9px;
    flex-wrap: wrap;
    align-items: center;
    margin: 16px 0;
}

.input,
select,
textarea {
    padding: 11px;
    border: 1px solid #cbd1d8;
    border-radius: 9px;
    background: #ffffff;
}

.grow {
    flex: 1;
}

.btn {
    border: 0;
    border-radius: 9px;
    padding: 11px 15px;
    font-weight: 700;
    text-decoration: none;
    cursor: pointer;
    display: inline-block;
}

.primary {
    background: var(--wine);
    color: #ffffff;
}

.light {
    background: #eceef1;
    color: #222222;
}

.danger {
    background: #fee2e2;
    color: #991b1b;
}

.form {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 14px;
}

.full {
    grid-column: 1 / -1;
}

.field label {
    display: block;
    font-weight: 700;
    font-size: 13px;
    margin-bottom: 5px;
}

.field input,
.field select,
.field textarea {
    width: 100%;
}

table {
    width: 100%;
    border-collapse: collapse;
}

th,
td {
    text-align: left;
    padding: 11px;
    border-bottom: 1px solid var(--line);
    font-size: 14px;
}

th {
    font-size: 12px;
    color: var(--muted);
    text-transform: uppercase;
}

.tag,
.status {
    display: inline-block;
    padding: 5px 9px;
    border-radius: 99px;
    background: #f4e5eb;
    color: var(--wine);
    font-weight: 700;
}

.pending {
    background: #fef3c7;
    color: #92400e;
}

.approved {
    background: #dcfce7;
    color: #166534;
}

.rejected {
    background: #fee2e2;
    color: #991b1b;
}

.flow {
    display: flex;
    gap: 18px;
    align-items: center;
    margin: 18px 0;
}

.factor {
    padding: 19px 28px;
    background: #f4f5f7;
    border-radius: 14px;
    text-align: center;
}

.factor b {
    display: block;
    font-size: 33px;
}

.competence-card {
    display: inline-block;
    margin: 10px 0 18px;
    padding: 10px 15px;
    background: #f4e5eb;
    border-radius: 10px;
    color: var(--wine);
    font-weight: 700;
}

.flash {
    padding: 12px;
    border-radius: 9px;
    margin-bottom: 13px;
}

.ok {
    background: #dcfce7;
    color: #166534;
}

.error {
    background: #fee2e2;
    color: #991b1b;
}

@media(max-width: 820px) {

    aside {
        position: static;
        width: auto;
    }

    .main {
        margin: 0;
    }

    .grid {
        grid-template-columns: 1fr 1fr;
    }

    .form {
        grid-template-columns: 1fr;
    }

    .full {
        grid-column: auto;
    }

    .content {
        padding: 18px;
    }

    .flow {
        align-items: stretch;
        flex-direction: column;
    }
}

</style>
"""


# ============================================================
# ESTRUTURA PRINCIPAL
# ============================================================

BASE = """
<!doctype html>

<html lang="pt-BR">

<head>

    <meta charset="utf-8">

    <meta
        name="viewport"
        content="width=device-width,initial-scale=1"
    >

    <title>V360 4.3</title>

""" + CSS + """

</head>

<body>

    <aside>

        <div class="logo">

            <b>V</b>360

            <small>
                VOLUME MARGEM 4.3
            </small>

        </div>

        <nav>

            <a href="/">
                ▦ Visão executiva
            </a>

            <a href="/catalog">
                ⌕ Catálogo por competência
            </a>

            <a href="/requests/new">
                ＋ Solicitar alteração
            </a>

            <a href="/requests">
                ✓ Análise do RH
            </a>

            <a href="/history">
                ☷ Histórico por competência
            </a>

            <a href="/setup">
                ↑ Carregar matriz
            </a>

            <a href="/governance">
                ⚙ Governança
            </a>

        </nav>

    </aside>

    <div class="main">

        <header class="top">

            <b>V360</b>

            <span>
                Fatores com vigência mensal
            </span>

        </header>

        <main class="content">

            {% with messages =
                get_flashed_messages(
                    with_categories=true
                )
            %}

                {% for category, message in messages %}

                    <div class="flash {{category}}">
                        {{message}}
                    </div>

                {% endfor %}

            {% endwith %}

            {{body|safe}}

        </main>

    </div>

</body>

</html>
"""


def page(body, **context):

    rendered_body = render_template_string(
        body,
        **context
    )

    return render_template_string(
        BASE,
        body=rendered_body
    )


@app.before_request
def boot():
    init()


# ============================================================
# CARREGAR MATRIZ
# ============================================================

@app.route(
    "/setup",
    methods=["GET", "POST"]
)
def setup():

    if request.method == "POST":

        file = request.files.get("file")

        if not file:

            flash(
                "Selecione uma planilha.",
                "error"
            )

            return redirect("/setup")

        path = ROOT / "matriz_carregada.xlsx"

        file.save(path)

        rows = extract(path)

        if not rows:

            flash(
                "Nenhum produto foi reconhecido.",
                "error"
            )

            return redirect("/setup")

        with cn() as c:

            c.execute(
                "DELETE FROM products"
            )

            c.execute(
                "DELETE FROM periods"
            )

            c.execute(
                "DELETE FROM requests"
            )

            c.executemany(
                """
                INSERT INTO products(
                    state,
                    category,
                    product,
                    base_factor,
                    matrix
                )
                VALUES(?,?,?,?,?)
                """,
                rows
            )

        log(
            "Carga de matriz",
            f"{len(rows)} registros"
        )

        flash(
            f"Matriz carregada: "
            f"{len(rows)} registros.",
            "ok"
        )

        return redirect("/")

    return page(
        """
        <div class="title">

            <h1>Carregar matriz</h1>

            <p>
                A carga cria a base por Estado,
                Categoria, Produto e Fator.
            </p>

        </div>

        <form
            method="post"
            enctype="multipart/form-data"
            class="panel"
        >

            <input
                class="input"
                type="file"
                name="file"
                accept=".xlsx"
                required
            >

            <button class="btn primary">
                Carregar
            </button>

        </form>
        """
    )


# ============================================================
# DASHBOARD
# ============================================================

@app.route("/")
def home():

    with cn() as c:

        indicators = c.execute(
            """
            SELECT
                COUNT(*) AS total,
                COUNT(DISTINCT product) AS products
            FROM products
            """
        ).fetchone()

        pending = c.execute(
            """
            SELECT COUNT(*) AS total
            FROM requests
            WHERE status = 'Pendente'
            """
        ).fetchone()["total"]

        scheduled = c.execute(
            """
            SELECT COUNT(*) AS total
            FROM periods
            WHERE competence >= ?
            """,
            (nowcomp(),)
        ).fetchone()["total"]

    return page(
        """
        <div class="title">

            <h1>Visão executiva</h1>

            <p>
                Alterações controladas
                por competência mensal.
            </p>

        </div>

        <div class="grid">

            <div class="card">

                <span>Produtos únicos</span>

                <strong>
                    {{indicators.products}}
                </strong>

            </div>

            <div class="card">

                <span>Registros</span>

                <strong>
                    {{indicators.total}}
                </strong>

            </div>

            <div class="card">

                <span>Pendentes RH</span>

                <strong>
                    {{pending}}
                </strong>

            </div>

            <div class="card">

                <span>Programadas</span>

                <strong>
                    {{scheduled}}
                </strong>

            </div>

        </div>

        {% if not indicators.total %}

            <div class="panel">

                <b>Base ainda não carregada.</b>

                Acesse “Carregar matriz”
                e selecione a planilha.

            </div>

        {% endif %}

        <div class="panel">

            <h3>Regra</h3>

            <p>
                O novo fator entra em vigor na
                competência aprovada e continua
                válido até uma alteração posterior.
                Meses anteriores permanecem intactos.
            </p>

        </div>
        """,
        indicators=indicators,
        pending=pending,
        scheduled=scheduled
    )


# ============================================================
# CATÁLOGO
# ============================================================

@app.route("/catalog")
def catalog():

    search = request.args.get(
        "q",
        ""
    ).strip()

    state = request.args.get(
        "uf",
        ""
    ).strip()

    category = request.args.get(
        "cat",
        ""
    ).strip()

    selected_factor = request.args.get(
        "factor",
        ""
    ).strip()

    competence = request.args.get(
        "competence",
        nowcomp()
    ).strip()

    if not competence:
        competence = nowcomp()

    with cn() as c:

        sql = """
            SELECT *
            FROM products
            WHERE 1 = 1
        """

        parameters = []

        if search:

            sql += " AND product LIKE ?"

            parameters.append(
                f"%{search}%"
            )

        if state:

            sql += " AND state = ?"
            parameters.append(state)

        if category:

            sql += " AND category = ?"
            parameters.append(category)

        sql += """
            ORDER BY product, state
            LIMIT 1600
        """

        products = c.execute(
            sql,
            parameters
        ).fetchall()

        rows = []

        for product in products:

            current_factor = factor(
                c,
                product["id"],
                competence
            )

            if (
                not selected_factor
                or current_factor
                == int(selected_factor)
            ):

                rows.append(
                    (
                        product,
                        current_factor
                    )
                )

        categories = c.execute(
            """
            SELECT DISTINCT category AS value
            FROM products
            ORDER BY category
            """
        ).fetchall()

    return page(
        """
        <div class="title">

            <h1>Catálogo por competência</h1>

            <p>
                Consulte o fator vigente
                em cada mês.
            </p>

        </div>

        <form class="toolbar">

            <input
                class="input grow"
                name="q"
                value="{{search}}"
                placeholder="Buscar produto"
            >

            <select name="uf">

                <option value="">
                    Estado
                </option>

                {% for value in ufs %}

                    <option
                        value="{{value}}"
                        {{'selected'
                        if state == value}}
                    >
                        {{value}}
                    </option>

                {% endfor %}

            </select>

            <select name="cat">

                <option value="">
                    Categoria
                </option>

                {% for item in categories %}

                    <option
                        value="{{item.value}}"
                        {{'selected'
                        if category == item.value}}
                    >
                        {{item.value}}
                    </option>

                {% endfor %}

            </select>

            <select name="factor">

                <option value="">
                    Fator
                </option>

                {% for value in [14,9,7,5,3,2] %}

                    <option
                        value="{{value}}"
                        {{'selected'
                        if selected_factor
                        == value|string}}
                    >
                        {{value}}
                    </option>

                {% endfor %}

            </select>

            <input
                class="input"
                type="month"
                name="competence"
                value="{{competence}}"
                required
            >

            <button class="btn primary">
                Filtrar
            </button>

            <a
                class="btn light"
                href="/catalog"
            >
                Limpar
            </a>

        </form>

        <div class="panel">

            <p>

                <b>{{rows|length}}</b>

                registros em

                <b>
                    {{competence|comp}}
                </b>

            </p>

            <table>

                <tr>
                    <th>Competência</th>
                    <th>Estado</th>
                    <th>Categoria</th>
                    <th>Produto</th>
                    <th>Fator</th>
                </tr>

                {% for product, current_factor in rows %}

                    <tr>

                        <td>
                            {{competence|comp}}
                        </td>

                        <td>
                            <span class="tag">
                                {{product.state}}
                            </span>
                        </td>

                        <td>
                            {{product.category}}
                        </td>

                        <td>
                            <b>
                                {{product.product}}
                            </b>
                        </td>

                        <td>
                            <b>
                                {{current_factor}}
                            </b>
                        </td>

                    </tr>

                {% endfor %}

            </table>

        </div>
        """,
        rows=rows,
        search=search,
        state=state,
        category=category,
        selected_factor=selected_factor,
        competence=competence,
        ufs=UFS,
        categories=categories
    )


# ============================================================
# NOVA SOLICITAÇÃO
# ============================================================

@app.route(
    "/requests/new",
    methods=["GET", "POST"]
)
def newreq():

    if request.method == "POST":

        product_id = int(
            request.form["product_id"]
        )

        competence = request.form.get(
            "competence",
            ""
        ).strip()

        new_factor = int(
            request.form["new_factor"]
        )

        area = request.form.get(
            "area",
            ""
        ).strip()

        other_area = request.form.get(
            "other_area",
            ""
        ).strip()

        if not competence:

            flash(
                "Informe a competência "
                "de início da alteração.",
                "error"
            )

            return redirect(
                "/requests/new"
            )

        if not re.fullmatch(
            r"\d{4}-\d{2}",
            competence
        ):

            flash(
                "A competência informada "
                "não é válida.",
                "error"
            )

            return redirect(
                "/requests/new"
            )

        if area == "Outros":

            if not other_area:

                flash(
                    "Informe a área no campo "
                    "Especificar Área.",
                    "error"
                )

                return redirect(
                    "/requests/new"
                )

            area = (
                f"Outros - {other_area}"
            )

        with cn() as c:

            product = c.execute(
                """
                SELECT *
                FROM products
                WHERE id = ?
                """,
                (product_id,)
            ).fetchone()

            if not product:

                flash(
                    "Produto não encontrado.",
                    "error"
                )

                return redirect(
                    "/requests/new"
                )

            current_factor = factor(
                c,
                product_id,
                competence
            )

            if current_factor == new_factor:

                flash(
                    "O novo fator deve ser "
                    "diferente do fator vigente "
                    "na competência selecionada.",
                    "error"
                )

                return redirect(
                    "/requests/new"
                )

            c.execute(
                """
                INSERT INTO requests(
                    created_at,
                    requester,
                    area,
                    product_id,
                    state,
                    category,
                    product,
                    old_factor,
                    new_factor,
                    competence,
                    justification
                )
                VALUES(
                    ?,?,?,?,?,?,?,?,?,?,?
                )
                """,
                (
                    datetime.now().isoformat(
                        timespec="seconds"
                    ),
                    request.form[
                        "requester"
                    ].strip(),
                    area,
                    product_id,
                    product["state"],
                    product["category"],
                    product["product"],
                    current_factor,
                    new_factor,
                    competence,
                    request.form[
                        "justification"
                    ].strip()
                )
            )

        log(
            "Solicitação",
            (
                f"{product['product']}/"
                f"{product['state']} "
                f"{competence}: "
                f"{current_factor}>"
                f"{new_factor}"
            )
        )

        flash(
            "Solicitação enviada "
            "com competência "
            f"{lab(competence)}.",
            "ok"
        )

        return redirect(
            "/requests"
        )

    with cn() as c:

        products = c.execute(
            """
            SELECT *
            FROM products
            ORDER BY product, state
            """
        ).fetchall()

    return page(
        """
        <div class="title">

            <h1>Solicitar alteração</h1>

            <p>
                Defina o mês de início
                do novo fator.
            </p>

        </div>

        <form
            method="post"
            class="panel form"
        >

            <div class="field">

                <label>Solicitante</label>

                <input
                    name="requester"
                    required
                >

            </div>

            <div class="field">

                <label>Área</label>

                <select
                    id="area"
                    name="area"
                    required
                    onchange="toggleOtherArea()"
                >

                    <option value="">
                        Selecione
                    </option>

                    <option value="Marketing">
                        Marketing
                    </option>

                    <option
                        value="Operações Comerciais"
                    >
                        Operações Comerciais
                    </option>

                    <option
                        value="Recursos Humanos"
                    >
                        Recursos Humanos
                    </option>

                    <option value="Outros">
                        Outros
                    </option>

                </select>

            </div>

            <div
                class="field"
                id="otherAreaContainer"
                style="display:none"
            >

                <label>
                    Especificar Área
                </label>

                <input
                    id="otherArea"
                    name="other_area"
                    placeholder="Digite a área"
                >

            </div>

            <div class="field full">

                <label>
                    Produto / Estado /
                    Categoria / Fator-base
                </label>

                <select
                    name="product_id"
                    required
                >

                    <option value="">
                        Selecione
                    </option>

                    {% for product in products %}

                        <option
                            value="{{product.id}}"
                        >
                            {{product.product}}
                            |
                            {{product.state}}
                            |
                            {{product.category}}
                            |
                            {{product.base_factor}}
                        </option>

                    {% endfor %}

                </select>

            </div>

            <div class="field">

                <label>Competência</label>

                <input
                    class="input"
                    type="month"
                    id="competence"
                    name="competence"
                    value="{{competence}}"
                    required
                >

            </div>

            <div class="field">

                <label>Novo fator</label>

                <select
                    name="new_factor"
                    required
                >

                    {% for value in [14,9,7,5,3,2] %}

                        <option value="{{value}}">
                            {{value}}
                        </option>

                    {% endfor %}

                </select>

            </div>

            <div class="field full">

                <label>Justificativa</label>

                <textarea
                    name="justification"
                    rows="4"
                    required
                ></textarea>

            </div>

            <button class="btn primary">
                Enviar
            </button>

        </form>

        <script>

        function toggleOtherArea() {

            const area =
                document.getElementById(
                    "area"
                );

            const container =
                document.getElementById(
                    "otherAreaContainer"
                );

            const field =
                document.getElementById(
                    "otherArea"
                );

            if (area.value === "Outros") {

                container.style.display =
                    "block";

                field.required = true;

            } else {

                container.style.display =
                    "none";

                field.required = false;

                field.value = "";

            }
        }

        document.addEventListener(
            "DOMContentLoaded",
            toggleOtherArea
        );

        </script>
        """,
        products=products,
        competence=nowcomp()
    )


# ============================================================
# FILA DE ANÁLISE
# ============================================================

@app.route("/requests")
def reqs():

    with cn() as c:

        requests_list = c.execute(
            """
            SELECT *
            FROM requests
            ORDER BY id DESC
            """
        ).fetchall()

    return page(
        """
        <div class="title">

            <h1>Análise do RH</h1>

            <p>
                Solicitações pendentes,
                aprovadas e rejeitadas.
            </p>

        </div>

        <div class="panel">

            <table>

                <tr>
                    <th>ID</th>
                    <th>Solicitante</th>
                    <th>Produto</th>
                    <th>Competência</th>
                    <th>Alteração</th>
                    <th>Status</th>
                    <th></th>
                </tr>

                {% for item in requests_list %}

                    <tr>

                        <td>
                            #{{item.id}}
                        </td>

                        <td>

                            {{item.requester}}

                            <br>

                            <small>
                                {{item.area}}
                            </small>

                        </td>

                        <td>

                            {{item.product}}

                            ·

                            {{item.state}}

                        </td>

                        <td>

                            <b>
                                {{item.competence|comp}}
                            </b>

                        </td>

                        <td>

                            {{item.old_factor}}

                            →

                            <b>
                                {{item.new_factor}}
                            </b>

                        </td>

                        <td>

                            <span class="
                                status
                                {{'pending'
                                if item.status
                                == 'Pendente'
                                else 'approved'
                                if item.status
                                == 'Aprovada'
                                else 'rejected'}}
                            ">

                                {{item.status}}

                            </span>

                        </td>

                        <td>

                            <a
                                class="btn light"
                                href="/requests/{{item.id}}"
                            >
                                Analisar
                            </a>

                        </td>

                    </tr>

                {% endfor %}

            </table>

        </div>
        """,
        requests_list=requests_list
    )


# ============================================================
# DETALHE E DECISÃO
# ============================================================

@app.route(
    "/requests/<int:request_id>",
    methods=["GET", "POST"]
)
def detail(request_id):

    with cn() as c:

        item = c.execute(
            """
            SELECT *
            FROM requests
            WHERE id = ?
            """,
            (request_id,)
        ).fetchone()

    if not item:
        return "Solicitação não encontrada.", 404

    if request.method == "POST":

        decision = request.form[
            "decision"
        ]

        decided_by = request.form[
            "decided_by"
        ].strip()

        reason = request.form.get(
            "reason",
            ""
        ).strip()

        decided_at = datetime.now().isoformat(
            timespec="seconds"
        )

        try:

            with cn() as c:

                item = c.execute(
                    """
                    SELECT *
                    FROM requests
                    WHERE id = ?
                    """,
                    (request_id,)
                ).fetchone()

                if item["status"] != "Pendente":

                    raise ValueError(
                        "Solicitação já analisada."
                    )

                competence = item["competence"]

                if not competence:

                    raise ValueError(
                        "A solicitação não possui "
                        "competência. Crie uma nova "
                        "solicitação antes de aprovar."
                    )

                if decision == "Aprovada":

                    c.execute(
                        """
                        INSERT OR REPLACE
                        INTO periods(
                            request_id,
                            product_id,
                            competence,
                            new_factor,
                            approved_at,
                            approved_by
                        )
                        VALUES(?,?,?,?,?,?)
                        """,
                        (
                            request_id,
                            item["product_id"],
                            competence,
                            item["new_factor"],
                            decided_at,
                            decided_by
                        )
                    )

                    action = "Aprovação"

                    details = (
                        f"#{request_id} "
                        f"{competence} "
                        f"{item['old_factor']}>"
                        f"{item['new_factor']}"
                    )

                else:

                    action = "Rejeição"

                    details = (
                        f"#{request_id}, "
                        f"sem alteração"
                    )

                c.execute(
                    """
                    UPDATE requests
                    SET
                        status = ?,
                        decided_at = ?,
                        decided_by = ?,
                        decision_reason = ?
                    WHERE id = ?
                    """,
                    (
                        decision,
                        decided_at,
                        decided_by,
                        reason,
                        request_id
                    )
                )

            log(
                action,
                details
            )

            if decision == "Aprovada":

                flash(
                    "Solicitação aprovada. "
                    "O novo fator foi programado "
                    f"para {lab(competence)}.",
                    "ok"
                )

            else:

                flash(
                    "Solicitação rejeitada. "
                    "O fator não foi alterado.",
                    "ok"
                )

        except Exception as error:

            flash(
                str(error),
                "error"
            )

        return redirect(
            "/requests"
        )

    return page(
        """
        <div class="title">

            <h1>
                Solicitação #{{item.id}}
            </h1>

            <p>
                <b>Solicitante:</b>
                {{item.requester}}
            </p>

            <p>
                <b>Área:</b>
                {{item.area}}
            </p>

            <div class="competence-card">

                Competência de início:

                {{item.competence|comp}}

            </div>

        </div>

        <div class="panel">

            <h3>

                {{item.product}}

                ·

                {{item.state}}

                ·

                {{item.category}}

            </h3>

            <div class="flow">

                <div class="factor">

                    Fator vigente

                    <b>
                        {{item.old_factor}}
                    </b>

                </div>

                <span>→</span>

                <div class="factor">

                    Fator solicitado

                    <b>
                        {{item.new_factor}}
                    </b>

                </div>

            </div>

            <p>

                <b>Competência:</b>

                {{item.competence|comp}}

            </p>

            <p>

                <b>Justificativa:</b>

                {{item.justification}}

            </p>

            {% if item.status == 'Pendente' %}

                <form
                    method="post"
                    class="form"
                >

                    <div class="field">

                        <label>
                            Responsável RH
                        </label>

                        <input
                            name="decided_by"
                            value="RH - Governança de Dados"
                            required
                        >

                    </div>

                    <div class="field full">

                        <label>Parecer</label>

                        <textarea
                            name="reason"
                            rows="3"
                        ></textarea>

                    </div>

                    <div class="toolbar full">

                        {% if item.competence %}

                            <button
                                class="btn primary"
                                name="decision"
                                value="Aprovada"
                            >
                                Aprovar para
                                {{item.competence|comp}}
                            </button>

                        {% else %}

                            <button
                                class="btn light"
                                type="button"
                                disabled
                            >
                                Competência não informada
                            </button>

                        {% endif %}

                        <button
                            class="btn danger"
                            name="decision"
                            value="Rejeitada"
                        >
                            Rejeitar
                        </button>

                    </div>

                </form>

            {% else %}

                <p>

                    <b>{{item.status}}</b>

                    por

                    {{item.decided_by}}

                </p>

                <p>

                    <b>Parecer:</b>

                    {{item.decision_reason}}

                </p>

            {% endif %}

        </div>
        """,
        item=item
    )


# ============================================================
# HISTÓRICO POR COMPETÊNCIA
# ============================================================

@app.route("/history")
def hist():

    with cn() as c:

        history = c.execute(
            """
            SELECT
                period.*,
                product.product,
                product.state,
                product.category,
                request.requester,
                request.area
            FROM periods AS period
            JOIN products AS product
                ON product.id =
                   period.product_id
            JOIN requests AS request
                ON request.id =
                   period.request_id
            ORDER BY
                period.competence DESC,
                period.id DESC
            """
        ).fetchall()

    return page(
        """
        <div class="title">

            <h1>
                Histórico por competência
            </h1>

            <p>
                Alterações efetivamente
                aprovadas pelo RH.
            </p>

        </div>

        <div class="panel">

            <table>

                <tr>
                    <th>Competência</th>
                    <th>Produto</th>
                    <th>Estado</th>
                    <th>Categoria</th>
                    <th>Fator</th>
                    <th>Solicitante</th>
                    <th>Aprovador</th>
                </tr>

                {% for item in history %}

                    <tr>

                        <td>
                            <b>
                                {{item.competence|comp}}
                            </b>
                        </td>

                        <td>
                            {{item.product}}
                        </td>

                        <td>
                            {{item.state}}
                        </td>

                        <td>
                            {{item.category}}
                        </td>

                        <td>
                            {{item.new_factor}}
                        </td>

                        <td>

                            {{item.requester}}

                            <br>

                            <small>
                                {{item.area}}
                            </small>

                        </td>

                        <td>
                            {{item.approved_by}}
                        </td>

                    </tr>

                {% endfor %}

            </table>

        </div>
        """,
        history=history
    )


# ============================================================
# GOVERNANÇA
# ============================================================

@app.route("/governance")
def gov():

    with cn() as c:

        audit_rows = c.execute(
            """
            SELECT *
            FROM audit
            ORDER BY id DESC
            """
        ).fetchall()

    return page(
        """
        <div class="title">

            <h1>Governança</h1>

        </div>

        <div class="panel">

            <p>
                ✓ A aprovação cria
                uma vigência mensal.
            </p>

            <p>
                ✓ O histórico anterior
                é preservado.
            </p>

            <p>
                ✓ A rejeição não altera
                a matriz.
            </p>

            <p>
                ✓ Solicitações sem competência
                não podem ser aprovadas.
            </p>

        </div>

        <div class="panel">

            <table>

                <tr>
                    <th>Data</th>
                    <th>Ação</th>
                    <th>Detalhes</th>
                </tr>

                {% for item in audit_rows %}

                    <tr>

                        <td>
                            {{item.created_at}}
                        </td>

                        <td>
                            {{item.action}}
                        </td>

                        <td>
                            {{item.details}}
                        </td>

                    </tr>

                {% endfor %}

            </table>

        </div>
        """,
        audit_rows=audit_rows
    )


# ============================================================
# INICIAR APLICAÇÃO
# ============================================================

if __name__ == "__main__":

    init()

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False
    )